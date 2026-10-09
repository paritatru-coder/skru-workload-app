import streamlit as st
import pandas as pd
import json
import io
import re
import time
import base64
import mimetypes
from datetime import datetime

import requests

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import fitz  # PyMuPDF (ใช้แปลงหน้า PDF เป็นรูปภาพเมื่อส่งไฟล์ตรงไม่ได้)
except ImportError:
    fitz = None

APP_VERSION = "v14"

CATEGORIES = [
    "1. ภาระงานสอน",
    "2. ภาระงานวิจัย / งานประพันธ์ / งานสร้างสรรค์",
    "3. ภาระงานบริการวิชาการ",
    "4. ภาระงานทำนุบำรุงศิลปวัฒนธรรม",
    "5. ภาระงานอื่น ๆ / งานสนับสนุน / คำสั่งเฉพาะกิจ",
    "6. งานประกันคุณภาพการศึกษา (QA)",
    "7. งานบริหาร / ตำแหน่งทางวิชาการ",
]
AUTO_CATEGORY = "ให้ AI ประเมินหมวดงานอัตโนมัติ"

DB_COLUMNS = ["email", "หมวดงาน", "รายการภาระงาน", "เลขคำสั่ง_อ้างอิง", "วันที่", "ภาระงาน_ชม", "วันที่บันทึก"]

# หมายเหตุ: Gemini 1.5 / 2.0 ถูกปิดแล้ว และ 2.5 Flash อยู่ในช่วงปิดให้บริการ
# gemini-flash-latest เป็น alias ที่ Google ชี้ไปรุ่น Flash ล่าสุดให้เอง จึงทนต่อการเปลี่ยนรุ่นได้ดีที่สุด
DEFAULT_MODEL = "gemini-flash-latest"
FALLBACK_MODELS = ["gemini-flash-latest", "gemini-3.5-flash", "gemini-3.1-flash-lite", "gemini-2.5-flash"]
MAX_INLINE_BYTES = 15 * 1024 * 1024  # เกินนี้จะแปลงเป็นรูปภาพทีละหน้าแทน

# ---------------------------------------------------------
# 1. Page Config & Custom Styling
# ---------------------------------------------------------
st.set_page_config(
    page_title=f"SKRU Workload AI - ระบบบันทึกและวิเคราะห์ภาระงาน มรภ.สงขลา ({APP_VERSION})",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
    .main-header { color: #1E3A8A; font-size: 2.2rem; font-weight: 700; margin-bottom: 0.2rem; }
    .sub-header { color: #4B5563; font-size: 1.05rem; margin-bottom: 1.2rem; }
    .success-alert {
        background-color: #DCFCE7; border: 2px solid #22C55E; color: #15803D;
        padding: 1rem; border-radius: 8px; font-weight: bold; font-size: 1.15rem; margin-bottom: 1rem;
    }
    .info-alert {
        background-color: #EFF6FF; border: 1px solid #93C5FD; color: #1E40AF;
        padding: 0.8rem; border-radius: 8px; margin-bottom: 1rem;
    }
    .formatted-preview {
        background-color: #F0F9FF; border: 1px solid #BAE6FD; border-radius: 8px;
        padding: 1rem; font-family: 'Sarabun', sans-serif; color: #0369A1; font-weight: 500;
        margin-top: 0.5rem; margin-bottom: 1rem;
    }
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# 2. Database Manager (Google Sheets / Local Session)
# ---------------------------------------------------------
conn = None
use_gsheets = False

try:
    from streamlit_gsheets import GSheetsConnection
    if hasattr(st, "secrets") and "connections" in st.secrets and "gsheets" in st.secrets["connections"]:
        conn = st.connection("gsheets", type=GSheetsConnection)
        use_gsheets = True
except Exception:
    use_gsheets = False


def ensure_schema(df):
    """ทำให้ DataFrame มีคอลัมน์ครบและชนิดข้อมูลถูกต้องเสมอ"""
    if df is None:
        df = pd.DataFrame(columns=DB_COLUMNS)
    df = df.copy()
    for c in DB_COLUMNS:
        if c not in df.columns:
            df[c] = 0.0 if c == "ภาระงาน_ชม" else ""
    df["ภาระงาน_ชม"] = pd.to_numeric(df["ภาระงาน_ชม"], errors="coerce").fillna(0.0)
    for c in DB_COLUMNS:
        if c != "ภาระงาน_ชม":
            df[c] = df[c].fillna("").astype(str)
    return df[DB_COLUMNS]


def load_all_data():
    if use_gsheets and conn:
        try:
            df = conn.read(ttl="0")
            if df is not None and not df.empty:
                df = df.dropna(how="all")
                return ensure_schema(df)
        except Exception:
            pass

    if "local_db" not in st.session_state:
        st.session_state["local_db"] = pd.DataFrame(columns=DB_COLUMNS)
    return ensure_schema(st.session_state["local_db"])


def save_all_data(full_df):
    full_df = ensure_schema(full_df)
    if use_gsheets and conn:
        try:
            conn.update(data=full_df)
            st.session_state["local_db"] = full_df
            return True, "Google Sheets"
        except Exception as e:
            st.session_state["local_db"] = full_df
            return False, str(e)
    st.session_state["local_db"] = full_df
    return True, "Local Session"


# ---------------------------------------------------------
# 3. PDF Text Layer Reader (ใช้เป็นข้อมูลเสริม / โหมดสำรอง)
# ---------------------------------------------------------
def extract_text_from_pdf_bytes(file_bytes):
    """อ่าน text layer จาก PDF (ถ้าเป็นภาพสแกนอาจได้ข้อความว่าง ซึ่งไม่ถือเป็นข้อผิดพลาด)"""
    text = ""
    if pdfplumber is not None:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text += t + "\n"
        except Exception:
            pass

    if not text.strip():
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
        except Exception:
            pass

    if not text.strip() and fitz is not None:
        try:
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                for page in doc:
                    t = page.get_text()
                    if t:
                        text += t + "\n"
        except Exception:
            pass

    return text.strip()


def render_pdf_pages_to_png(file_bytes, max_pages=8, zoom=2.0):
    """แปลงหน้า PDF เป็นรูป PNG (ต้องมี PyMuPDF)"""
    if fitz is None:
        raise RuntimeError("ไม่พบไลบรารี PyMuPDF (pip install PyMuPDF)")
    images = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for i, page in enumerate(doc):
            if i >= max_pages:
                break
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
            images.append(pix.tobytes("png"))
    if not images:
        raise RuntimeError("PDF ไม่มีหน้าที่แปลงเป็นภาพได้")
    return images


# ---------------------------------------------------------
# 4. Gemini Vision Engine (REST API - ไม่ต้องพึ่ง SDK)
# ---------------------------------------------------------
class GeminiError(Exception):
    pass


class GeminiAuthError(GeminiError):
    pass


class GeminiRequestError(GeminiError):
    pass


class GeminiUnavailable(GeminiError):
    pass


RUBRIC_TEXT = """
เกณฑ์ภาระงาน มรภ.สงขลา (สรุป):
- ผลงานวิชาการ/วิจัย: Proceedings ระดับชาติ 10 ชม.; Proceedings ระดับนานาชาติ หรือวารสารที่สภามหาวิทยาลัยอนุมัติ 12; TCI กลุ่ม 2 = 15; TCI กลุ่ม 1 = 19; บทความที่ประกาศใน ก.พ.อ. = 21; บทความได้รับรางวัลระดับชาติ/นานาชาติ = 14/20; ผลงานวิจัยบูรณาการ = 4. ทุนวิจัย (ชม./สัปดาห์): คณะ 6, มหาวิทยาลัย 8, ภายนอก 10
- ตำรา/หนังสือ ไม่เกิน 12; เอกสารคำสอน ไม่เกิน 9; เอกสารประกอบการสอน ไม่เกิน 6; สื่อประกอบการสอน ไม่เกิน 6
- งานสร้างสรรค์/สิ่งประดิษฐ์: ระดับชาติไม่รางวัล ไม่เกิน 5; ระดับชาติรางวัล ไม่เกิน 10; ความร่วมมือระหว่างประเทศ ไม่เกิน 12; อาเซียนไม่รางวัล ไม่เกิน 15; อาเซียนรางวัล ไม่เกิน 18; นานาชาติไม่รางวัล ไม่เกิน 21; นานาชาติรางวัล ไม่เกิน 24
- บริการวิชาการ: ผู้รับผิดชอบหลัก 5; ประธานฝ่าย 2; กรรมการ 1; เลขานุการ 1.5; วิทยากร 0.5 ชม./1 ชม.; ผลงานบูรณาการ 4
- บริการวิชาการที่มีรายได้: <25,000 = 2; 25,001-50,000 = 3; 50,000-75,000 = 4; 75,000-100,000 = 5; >100,000 = 6
- ทำนุบำรุงศิลปวัฒนธรรม (ผู้จัด): ประธาน 1 ชม./วัน; กรรมการ 0.5 ชม./วัน; เลขานุการ 0.75 ชม./วัน; เข้าร่วมโครงการ 0.5 ชม./โครงการ
- การอ่านผลงาน: ระดับอุดมศึกษา 2 ชม./ผลงาน; เพื่อเผยแพร่/ตำแหน่งวิชาการ 1 ชม./ผลงาน; บรรณาธิการวารสาร 3; กองบรรณาธิการ 1; กรรมการวิชาการ/ผู้เชี่ยวชาญภายนอก 1 ชม./คำสั่ง
- ที่ปรึกษาโครงงานปริญญาตรี: ที่ปรึกษา 1 ชม./คน; ที่ปรึกษาร่วม 0.75 ชม./คน (รวมไม่เกิน 10); กรรมการสอบ หัวข้อ 0.15 เค้าโครง 0.3 โครงการ 0.5 ชม./เรื่อง; ที่ปรึกษาวิทยานิพนธ์ 3 ชม./สัปดาห์; ที่ปรึกษาร่วม/ค้นคว้าอิสระ 2 ชม./สัปดาห์
- พัฒนานักศึกษา: อาจารย์ที่ปรึกษา 1-30 คน 2 ชม./หมู่เรียน, เกิน 30 คน 3 ชม./หมู่เรียน; ประธานที่ปรึกษาโครงการ 3; กรรมการที่ปรึกษา 1.5; ผู้จัดการทีม 1 ชม./คำสั่ง; ผู้ฝึกสอนกีฬา 3 ชม./คำสั่ง; ประธาน 1 / กรรมการ 0.5 / เลขานุการ 0.75 ชม./คำสั่ง
- คำสั่งเฉพาะกิจ (ต่อคำสั่ง): ประธาน 1 / กรรมการ 0.5 / เลขานุการ 0.75; โครงการเกิน 1 ภาคการศึกษา: 3 / 1 / 2; สอบข้อเท็จจริงหรือวินัย: 5 / 2 / 3
- พัฒนาตนเอง: เข้าร่วมตามแผนพัฒนารายบุคคล ในประเทศ 1 ชม./วันทำการ, ต่างประเทศ 1.5 ชม./วันทำการ; ผลงานสืบเนื่องจากการพัฒนาตนเอง 4 ชม./ผลงาน
- งานบริหาร: อธิการบดี 35; รองอธิการบดี/คณบดี 30; ผู้ช่วยอธิการบดี/รองคณบดี 25; ประธานหลักสูตร 15; เลขานุการหลักสูตร/ผู้ควบคุมหอพัก 10; กรรมการบริหารหลักสูตร 8; กรรมการสโมสรอาจารย์ 2
- งานพัสดุ: ให้คิดตามบทบาท ประธาน/กรรมการ/เลขานุการ ต่อคำสั่ง ตามวงเงินในเอกสาร
"""


def build_prompt(category_hint, local_text):
    cat_list = "\n".join(CATEGORIES)
    hint_line = ""
    if category_hint and category_hint != AUTO_CATEGORY:
        hint_line = "ผู้ใช้ระบุหมวดงานเบื้องต้นไว้แล้วว่า: " + category_hint + " (ให้ใช้หมวดนี้)\n"

    prompt = (
        "คุณคือผู้ช่วยตรวจเอกสารภาระงานสายวิชาการของมหาวิทยาลัยราชภัฏสงขลา "
        "อ่านเอกสารที่แนบมา (อาจเป็นภาพสแกนภาษาไทย) อย่างละเอียดทุกหน้า แล้วสกัดข้อมูลจริงจากเนื้อหาในเอกสาร\n\n"
        + hint_line
        + "กติกาสำคัญ:\n"
        "1. ห้ามใช้ชื่อไฟล์ (.pdf/.png/.jpg) เป็นข้อมูลใดๆ ทั้งสิ้น และห้ามใส่ชื่อ-นามสกุลผู้รับเอกสารลงในช่องใดๆ\n"
        "2. ถ้าเอกสารมีหลายหน้าที่เป็นใบเดียวกันคนละภาษา (เช่น ไทย/อังกฤษ) ให้นับเป็น 1 รายการ ห้ามนับซ้ำ\n"
        "3. doc_type เลือกหนึ่งค่า: certificate_training (ใบประกาศนียบัตร/เกียรติบัตรผ่านการอบรม เรียนออนไลน์ MOOC สัมมนา), "
        "command (คำสั่ง/ประกาศแต่งตั้ง), creative_research (ผลงานวิจัย/บทความ/ประพันธ์/งานสร้างสรรค์/รางวัลของผลงาน), "
        "service (บริการวิชาการ/วิทยากร), other\n"
        "4. ถ้า doc_type = certificate_training ให้กรอก: course_name = ชื่อหลักสูตรเท่านั้น (ไม่รวมคำว่า 'รายวิชา' 'หลักสูตร' และไม่รวมชื่อผู้รับ), "
        "issuer = หน่วยงาน/โครงการที่ออกใบประกาศ (เช่น Thai MOOC กระทรวง อว.), cert_code = รหัส/เลขที่ใบประกาศที่พิมพ์ในเอกสาร (เช่นรหัสใต้ QR code), "
        "course_hours = จำนวนชั่วโมงของหลักสูตรที่ระบุในเอกสาร (ตัวเลข), date = วันที่ออกใบประกาศ, และตั้ง hours = 0 (ระบบคำนวณเอง)\n"
        "5. title = ชื่อผลงาน/โครงการ/บทบาทที่ระบุไว้จริงในเอกสาร; venue = เวที/สถานที่/หน่วยงานผู้จัดจริง\n"
        "6. date = วันที่จริงในเอกสาร (รูปแบบไทย พ.ศ. เช่น 16 มิถุนายน 2569; หลายวันให้คั่นด้วยจุลภาค); ref = เลขที่หนังสือ/คำสั่ง/ประกาศอ้างอิงจริง\n"
        "7. ถ้าไม่พบข้อมูลช่องใดให้ใส่สตริงว่าง \"\" หรือ 0 ห้ามเดา ห้ามแต่งขึ้นเอง\n"
        "8. level = ระดับของผลงาน เลือกจาก: ระดับชาติ, อาเซียน, นานาชาติ, ความร่วมมือระหว่างประเทศ, ไม่เกี่ยวข้อง/ไม่ระบุ\n"
        "9. has_award = true เฉพาะเมื่อ 'ผลงาน' ได้รับรางวัล/ผลประเมินดีเยี่ยม/Excellent จริง (คำว่า 'Awarded by/to' ในใบประกาศอบรมไม่นับเป็นรางวัล)\n"
        "10. สำหรับ doc_type อื่น: hours = จำนวนภาระงานตามเกณฑ์ด้านล่าง; formula = สูตรสั้นๆ ในวงเล็บ เช่น (3ชมx5วันx0.5)\n"
        "11. category ต้องเป็นข้อความหนึ่งในรายการต่อไปนี้ทุกตัวอักษร:\n" + cat_list + "\n"
        "12. raw_text = ข้อความสำคัญที่อ่านได้ (ไม่เกิน 1500 ตัวอักษร); notes = สิ่งที่ผู้ใช้ควรตรวจสอบซ้ำ (ถ้ามี)\n\n"
        + RUBRIC_TEXT
        + "\nตอบกลับเป็น JSON เท่านั้น ห้ามมีข้อความอื่นนอก JSON ตามโครงสร้าง:\n"
        '{"doc_type": "", "category": "", "title": "", "venue": "", "date": "", "ref": "", "level": "", '
        '"has_award": false, "course_name": "", "issuer": "", "cert_code": "", "course_hours": 0, '
        '"formula": "", "hours": 0, "raw_text": "", "notes": ""}\n'
    )
    if local_text:
        prompt += "\nข้อความที่ผู้ใช้พิมพ์/ดึงมา:\n" + local_text[:6000] + "\n"
    return prompt


def extract_json(text):
    t = (text or "").strip()
    t = re.sub(r"^```(?:json)?\s*", "", t, flags=re.I)
    t = re.sub(r"\s*```$", "", t)
    try:
        return json.loads(t)
    except Exception:
        pass
    m = re.search(r"\{.*\}", t, re.S)
    if m:
        try:
            return json.loads(m.group(0))
        except Exception:
            pass
    raise GeminiRequestError("Gemini ตอบกลับในรูปแบบที่อ่านเป็น JSON ไม่ได้")


def inline_part(data_bytes, mime):
    return {"inline_data": {"mime_type": mime, "data": base64.b64encode(data_bytes).decode("ascii")}}


def list_available_models(api_key):
    """ถาม Google ว่าบัญชีนี้ใช้รุ่นไหนได้บ้าง (เฉพาะ Flash ที่รองรับ generateContent)"""
    try:
        resp = requests.get(
            "https://generativelanguage.googleapis.com/v1beta/models?pageSize=200",
            headers={"x-goog-api-key": api_key}, timeout=30,
        )
        if resp.status_code != 200:
            return []
        names = []
        for m in resp.json().get("models", []):
            name = m.get("name", "").replace("models/", "")
            methods = m.get("supportedGenerationMethods", [])
            if "flash" in name and "generateContent" in methods:
                if not any(x in name for x in ("image", "tts", "live", "embedding", "audio", "omni", "robotics")):
                    names.append(name)
        return sorted(set(names), reverse=True)
    except Exception:
        return []


def call_gemini(api_key, models, parts):
    """เรียก Gemini ตามลำดับรุ่น คืนค่า (dict, model_ที่ใช้)"""
    body = {
        "contents": [{"parts": parts}],
        "generationConfig": {
            "temperature": 0,
            "responseMimeType": "application/json",
            "maxOutputTokens": 8192,
            "thinkingConfig": {"thinkingLevel": "low"},  # ลดเวลาคิด -> เร็วขึ้น (ถ้ารุ่นไม่รองรับ จะตัดออกแล้วลองใหม่อัตโนมัติ)
        },
    }
    headers = {"Content-Type": "application/json", "x-goog-api-key": api_key}
    last_err = "ไม่ทราบสาเหตุ"
    request_error = None
    not_found = []
    other_failed = False

    for model in models:
        url = "https://generativelanguage.googleapis.com/v1beta/models/" + model + ":generateContent"
        for attempt in range(2):
            try:
                resp = requests.post(url, headers=headers, json=body, timeout=120)
            except requests.RequestException as e:
                last_err = "เชื่อมต่อ Gemini ไม่ได้: " + str(e)
                time.sleep(1.0)
                continue

            code = resp.status_code
            if code == 200:
                try:
                    payload = resp.json()
                    cands = payload.get("candidates") or []
                    if not cands:
                        raise GeminiRequestError("Gemini ไม่ส่งผลลัพธ์กลับมา (อาจถูกกรองเนื้อหา)")
                    texts = [p.get("text", "") for p in cands[0].get("content", {}).get("parts", [])]
                    return extract_json("".join(texts)), model
                except GeminiRequestError as e:
                    request_error = e
                    break
                except Exception as e:
                    request_error = GeminiRequestError("อ่านผลลัพธ์จาก Gemini ไม่สำเร็จ: " + str(e))
                    break

            try:
                err_msg = resp.json().get("error", {}).get("message", resp.text[:200])
            except Exception:
                err_msg = resp.text[:200]

            if code == 400 and "thinking" in err_msg.lower() and "thinkingConfig" in body["generationConfig"]:
                del body["generationConfig"]["thinkingConfig"]
                continue
            if code in (401, 403) or (code == 400 and ("API key" in err_msg or "API_KEY" in err_msg)):
                raise GeminiAuthError("API Key ไม่ถูกต้องหรือไม่มีสิทธิ์ใช้งาน: " + err_msg)
            if code == 404:
                not_found.append(model)
                last_err = "ไม่พบรุ่น " + ", ".join(not_found) + " (อาจถูกปิดให้บริการแล้ว)"
                break  # ลองรุ่นถัดไป
            if code == 400:
                request_error = GeminiRequestError("Gemini ปฏิเสธคำขอ: " + err_msg)
                break
            if code in (429, 500, 502, 503, 504):
                other_failed = True
                last_err = "Gemini ไม่ว่าง/เกินโควตา (" + str(code) + "): " + err_msg
                time.sleep(2.0 * (attempt + 1))
                continue
            last_err = "Gemini ตอบกลับสถานะ " + str(code) + ": " + err_msg
            break

        if request_error is not None:
            raise request_error

    # ทุกรุ่นที่ลองไม่พบ (404) -> ถามรายชื่อรุ่นที่บัญชีนี้ใช้ได้จริงแล้วลองรุ่นที่ยังไม่เคยลอง
    if not_found and not other_failed and len(not_found) == len(models):
        extra = [m for m in list_available_models(api_key) if m not in not_found][:3]
        if extra:
            return call_gemini(api_key, extra, parts)
        last_err += " และดึงรายชื่อรุ่นที่ใช้ได้ไม่สำเร็จ กรุณาใส่ชื่อรุ่นในช่อง 'รุ่น Gemini' ด้านซ้ายเอง (ดูรายชื่อที่ ai.google.dev/gemini-api/docs/models)"

    raise GeminiUnavailable(last_err)


def analyze_with_gemini(api_key, models, file_bytes, mime, local_text, category_hint, typed_text=""):
    """ส่งไฟล์ (PDF ตรง / รูปภาพ / ข้อความ) เข้า Gemini Vision"""
    # ไม่ส่ง text layer ไปกับไฟล์: PDF/ภาพถูกส่งเข้า Gemini โดยตรงอยู่แล้ว และ text layer ภาษาไทยมักสระ/วรรณยุกต์หาย
    prompt = build_prompt(category_hint, "")

    if typed_text:
        parts = [{"text": prompt + "\nข้อความที่ผู้ใช้พิมพ์เอง:\n" + typed_text[:8000]}]
        return call_gemini(api_key, models, parts)

    if mime == "application/pdf":
        first_error = None
        if len(file_bytes) <= MAX_INLINE_BYTES:
            try:
                parts = [{"text": prompt}, inline_part(file_bytes, "application/pdf")]
                return call_gemini(api_key, models, parts)
            except GeminiRequestError as e:
                first_error = e  # ลองแปลงเป็นรูปภาพต่อ
        try:
            images = render_pdf_pages_to_png(file_bytes)
        except Exception as e:
            if first_error is not None:
                raise first_error
            raise GeminiRequestError("ไฟล์ PDF ใหญ่เกินไปและแปลงเป็นภาพไม่ได้: " + str(e))
        parts = [{"text": prompt}] + [inline_part(img, "image/png") for img in images]
        return call_gemini(api_key, models, parts)

    parts = [{"text": prompt}, inline_part(file_bytes, mime)]
    return call_gemini(api_key, models, parts)


# ---------------------------------------------------------
# 5. Business Rules & Normalization
# ---------------------------------------------------------
AWARD_RE = re.compile(r"ดีเยี่ยม|รางวัล|\bexcellent\b|\baward(?:s)?\b", re.I)  # "Awarded by/to" ไม่ถูกนับ
TRAINING_RE = re.compile(
    r"certificate of completion|thai\s*mooc|ผ่านเกณฑ์หลักสูตร|ผ่านการอบรม|ได้เข้ารับการอบรม|completion and fulfil+ment", re.I
)
THAI_MARKS = set("่้๊๋ิีึืุูัํ็์ำ")
INTL_KEYWORDS = ["นานาชาติ", "international", "intl"]
BAD_VALUES = {"none", "null", "n/a", "-", "ไม่ระบุ", "ไม่พบ", "ไม่มี"}
FILE_EXT_RE = re.compile(r"\.(pdf|png|jpe?g|webp)\b", re.I)


def clean_field(value, filename=""):
    """ล้างค่าที่ว่าง/ไม่ระบุ และตัดทิ้งหากเป็นชื่อไฟล์ (ห้ามนำชื่อไฟล์มาใช้)"""
    if value is None:
        return ""
    s = str(value).strip()
    if s.lower() in BAD_VALUES:
        return ""
    if FILE_EXT_RE.search(s):
        return ""
    if filename:
        stem = filename.rsplit(".", 1)[0].strip().lower()
        if s.lower() == filename.strip().lower() or (stem and s.lower() == stem):
            return ""
    return s


def normalize_category(value, hint=""):
    if hint and hint != AUTO_CATEGORY and hint in CATEGORIES:
        return hint
    v = (value or "").strip()
    if v in CATEGORIES:
        return v
    m = re.match(r"\s*([1-7])\b", v)
    if m:
        return CATEGORIES[int(m.group(1)) - 1]
    return CATEGORIES[1]


def to_float(value, default=0.0):
    try:
        f = float(value)
        return f if f >= 0 else default
    except Exception:
        return default


def contains_any(text, keywords):
    low = (text or "").lower()
    return any(k in low for k in keywords)


def compose_formal_text(title, venue, date, ref, formula, hours):
    head = " ".join(p.strip() for p in [title, venue, date] if p and p.strip())
    if ref and ref.strip():
        head += " (" + ref.strip() + ")"
    head = head.strip()
    if formula and formula.strip():
        return f"{head} = {formula.strip()} = {hours:.1f} ชม."
    return f"{head} = {hours:.1f} ชม."


def thai_text_looks_garbled(text):
    """text layer ไทยที่สระ/วรรณยุกต์หาย (พบบ่อยใน PDF จากระบบออกใบประกาศ)"""
    thai = [c for c in (text or "") if "\u0e01" <= c <= "\u0e5b"]
    if len(thai) < 30:
        return False
    consonants = sum(1 for c in thai if "\u0e01" <= c <= "\u0e2e")
    marks = sum(1 for c in thai if c in THAI_MARKS)
    return consonants > 0 and (marks / consonants) < 0.24


def finalize_result(raw, evidence_text, filename, category_hint, training_mode="cert"):
    """ปรับผลลัพธ์ให้เป็นมาตรฐาน + บังคับใช้กฎ (ใบประกาศอบรม -> หมวด 5, งานนานาชาติรางวัล -> 24)"""
    category = normalize_category(raw.get("category", ""), category_hint)
    hint_locked = bool(category_hint and category_hint != AUTO_CATEGORY and category_hint in CATEGORIES)
    title = clean_field(raw.get("title"), filename)
    venue = clean_field(raw.get("venue"), filename)
    date_str = clean_field(raw.get("date"), filename)
    ref_code = clean_field(raw.get("ref"), filename)
    formula = clean_field(raw.get("formula"), filename)
    hours = to_float(raw.get("hours"), 0.0)
    level = str(raw.get("level", "") or "")
    notes = clean_field(raw.get("notes"), filename)
    doc_type = str(raw.get("doc_type", "") or "").strip()

    all_text = " ".join([evidence_text or "", str(raw.get("raw_text", "") or ""), title, venue, ref_code, formula, level, notes,
                         str(raw.get("course_name", "") or ""), str(raw.get("issuer", "") or "")])

    is_intl = contains_any(all_text, INTL_KEYWORDS) or "นานาชาติ" in level
    has_award = bool(raw.get("has_award")) or bool(AWARD_RE.search(all_text))

    is_training = (doc_type == "certificate_training") or (
        doc_type in ("", "other") and bool(TRAINING_RE.search(all_text)) and not (is_intl and has_award)
    )

    rule_applied = False
    if is_training:
        # ใบประกาศ/เกียรติบัตรอบรม -> หมวด 5 เสมอ (ยกเว้นผู้ใช้เลือกหมวดเอง)
        if not hint_locked:
            category = CATEGORIES[4]
        course = clean_field(raw.get("course_name"), filename)
        course = re.sub(r"^(หลักสูตร|รายวิชา)\s*", "", course).strip()
        issuer = clean_field(raw.get("issuer"), filename)
        code = clean_field(raw.get("cert_code"), filename)
        course_hours = to_float(raw.get("course_hours"), 0.0)
        if course:
            title = "ผ่านการอบรมหลักสูตร '" + course + "'"
        if issuer:
            venue = issuer
        if code:
            ref_code = "ใบประกาศนียบัตร รหัส " + code
        if training_mode == "rubric":
            hours = 1.0
            formula = "(อบรม 1 วันทำการ x 1 ชม. ตามเกณฑ์งานพัฒนาตนเอง = 1 ภาระงาน)"
        elif course_hours > 0:
            hours = course_hours
            formula = f"(อบรมหลักสูตร {course_hours:g} ชั่วโมง = {course_hours:g} ภาระงาน)"
    elif category == CATEGORIES[1] and is_intl and has_award:
        # กฎเฉพาะ: งานสร้างสรรค์/วิจัยนานาชาติ ที่พบ ดีเยี่ยม/Excellent/รางวัล/Award = 24 ภาระงานเสมอ
        hours = 24.0
        formula = "(นับสิทธิ์เผยแพร่นานาชาติ ที่ได้รับรางวัล/ระดับดีเยี่ยม = 24 ภาระงาน)"
        rule_applied = True

    return {
        "category": category,
        "title": title,
        "venue": venue,
        "date": date_str,
        "ref": ref_code,
        "formula": formula,
        "hours": hours,
        "formal_text": compose_formal_text(title, venue, date_str, ref_code, formula, hours),
        "notes": notes,
        "rule_applied": rule_applied,
        "is_training": is_training,
    }


# ---------------------------------------------------------
# 6. Local Fallback Parser (ใช้เมื่อไม่มี/เรียก Gemini ไม่ได้)
# ---------------------------------------------------------
_MONTHS_FULL = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
                "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
_MONTHS_ABBR = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
_month_alts = _MONTHS_FULL + _MONTHS_ABBR + [a.rstrip(".") for a in _MONTHS_ABBR]
_month_alts = sorted(set(_month_alts), key=len, reverse=True)
DATE_RE = re.compile(
    r"\d{1,2}(?:\s*(?:,|-|–|และ)\s*\d{1,2})*\s*(?:" + "|".join(re.escape(m) for m in _month_alts) + r")\s*(?:25|26)\d{2}"
)
REF_RES = [
    re.compile(r"(?:คำสั่ง|ประกาศ)[^\n]{0,60}?ที่\s*[0-9ก-ฮA-Za-z\.\-]+\s*/\s*\d{4}"),
    re.compile(r"ที่\s+[ก-ฮA-Za-z]{1,8}[0-9\.]*\s*/\s*[0-9]{1,6}"),
]
TITLE_RE = re.compile(r"เรื่อง\s+([^\n]{5,200})")
VENUE_RE = re.compile(r"(ณ\s+[^\n]{3,100})")


def guess_category_local(text):
    t = text or ""
    if contains_any(t, ["วิทยากร", "บริการวิชาการ", "โครงการอบรม"]):
        return CATEGORIES[2]
    if contains_any(t, ["ประพันธ์", "สร้างสรรค์", "วิจัย", "proceedings", "journal", "บทความ", "นานาชาติ"]):
        return CATEGORIES[1]
    if contains_any(t, ["ทำนุบำรุง", "ศิลปวัฒนธรรม"]):
        return CATEGORIES[3]
    if contains_any(t, ["ประกันคุณภาพ", "qa"]):
        return CATEGORIES[5]
    return CATEGORIES[4]


def parse_text_local(text, filename, category_hint, training_mode="cert"):
    t = text or ""
    raw = {"category": guess_category_local(t), "title": "", "venue": "", "date": "", "ref": "",
           "formula": "", "hours": 0.0, "level": "", "has_award": False, "raw_text": t[:2500], "notes": ""}

    m = TITLE_RE.search(t)
    if m:
        raw["title"] = m.group(1).strip()
    m = VENUE_RE.search(t)
    if m:
        raw["venue"] = m.group(1).strip()
    m = DATE_RE.search(t)
    if m:
        raw["date"] = re.sub(r"\s+", " ", m.group(0)).strip()
    for rx in REF_RES:
        m = rx.search(t)
        if m:
            raw["ref"] = re.sub(r"\s+", " ", m.group(0)).strip()
            break
    return finalize_result(raw, t, filename, category_hint, training_mode)


def blank_result(category_hint):
    cat = normalize_category("", category_hint)
    return {"category": cat, "title": "", "venue": "", "date": "", "ref": "", "formula": "",
            "hours": 0.0, "formal_text": "", "notes": "", "rule_applied": False, "is_training": False}


def guess_mime(uploaded):
    mime = getattr(uploaded, "type", None)
    if mime and mime != "application/octet-stream":
        return mime
    guess, _ = mimetypes.guess_type(getattr(uploaded, "name", "") or "")
    return guess or "image/jpeg"


def sniff_mime(data):
    if data[:4] == b"%PDF":
        return "application/pdf"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return ""


def fetch_drive_file(url):
    """ดาวน์โหลดไฟล์จากลิงก์ Google Drive ที่แชร์แบบ 'ทุกคนที่มีลิงก์' (ใช้แทนการอัปโหลดจากมือถือที่ไม่ผ่าน)"""
    m = re.search(r"/d/([A-Za-z0-9_-]{10,})", url) or re.search(r"[?&]id=([A-Za-z0-9_-]{10,})", url)
    if not m:
        raise ValueError("ลิงก์ Google Drive ไม่ถูกต้อง (ควรเป็นรูปแบบ drive.google.com/file/d/.../view)")
    try:
        resp = requests.get(
            "https://drive.usercontent.google.com/download",
            params={"id": m.group(1), "export": "download", "confirm": "t"},
            timeout=60,
        )
    except requests.RequestException as e:
        raise ValueError("ดาวน์โหลดจาก Google Drive ไม่ได้: " + str(e))
    data = resp.content
    if len(data) > 25 * 1024 * 1024:
        raise ValueError("ไฟล์ใหญ่เกิน 25 MB")
    mime = sniff_mime(data)
    if resp.status_code != 200 or not mime:
        raise ValueError("เปิดไฟล์จากลิงก์ไม่ได้ — ตรวจว่าตั้งการแชร์เป็น 'ทุกคนที่มีลิงก์' (Anyone with the link) และไฟล์เป็น PDF/PNG/JPG")
    return data, mime


@st.cache_data(show_spinner=False, ttl=3600, max_entries=30)
def cached_analyze(file_bytes, mime, category_hint, typed_text, models_tuple, _api_key):
    """แคชผลต่อไฟล์เดียวกัน -> กดซ้ำไม่เสียเวลา/โควตา และได้ผลเหมือนเดิมทุกครั้ง"""
    return analyze_with_gemini(_api_key, list(models_tuple), file_bytes, mime, "", category_hint, typed_text)


# ---------------------------------------------------------
# 7. Sidebar Configuration
# ---------------------------------------------------------
with st.sidebar:
    st.title("👤 ระบุตัวตน & ตั้งค่า")

    user_email = st.text_input(
        "📧 อีเมลบุคลากร (สำหรับซิงค์ข้อมูล):",
        value=st.session_state.get("user_email_saved", ""),
        placeholder="เช่น instructor@skru.ac.th",
        help="กรอกอีเมลบุคลากรของคุณ เพื่อซิงค์คลังภาระงานส่วนตัวจากมือถือ คอมพิวเตอร์บ้าน และที่ทำงาน",
        key="sidebar_email_input",
    )
    if user_email:
        st.session_state["user_email_saved"] = user_email.strip()

    track_type = st.selectbox(
        "🎯 กลุ่มการประเมินสายวิชาการ:",
        ["กลุ่มเน้นการสอน (สอน 30%, วิจัย 15-20%, บริการ 8-15%, ศิลปะ 2-5%, อื่นๆ 5-15%, QA 1%)",
         "กลุ่มเน้นการวิจัย (สอน 25%, วิจัย 15-25%, บริการ 8-15%, ศิลปะ 2-5%, อื่นๆ 5-15%, QA 1%)"],
    )

    role = st.selectbox(
        "🏛️ ตำแหน่งบริหาร/วิชาการ:",
        [
            "อาจารย์สายสอน (ไม่ดำรงตำแหน่งบริหาร)",
            "ประธานหลักสูตร / หัวหน้าสาขา (15 ภาระงาน)",
            "เลขานุการหลักสูตร (10 ภาระงาน)",
            "ผู้ช่วยอธิการบดี / รองคณบดี (25 ภาระงาน)",
            "คณบดี / หัวหน้าหน่วยงาน (30 ภาระงาน)",
        ],
    )

    st.divider()

    api_key_env = ""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            api_key_env = str(st.secrets["GEMINI_API_KEY"]).strip()
    except Exception:
        api_key_env = ""

    user_api_key = st.text_input(
        "🔑 Gemini API Key (ถ้าระบุเอง):",
        type="password",
        value="",
        placeholder="ใช้ค่าจาก Secrets อยู่" if api_key_env else "วาง API Key ที่นี่",
    )
    active_api_key = (user_api_key or "").strip() or api_key_env

    model_choice = st.text_input("🧠 รุ่น Gemini:", value=DEFAULT_MODEL, help="หากรุ่นนี้ใช้ไม่ได้ ระบบจะลองรุ่นสำรองให้อัตโนมัติ")
    model_list = [model_choice.strip()] if model_choice.strip() else []
    for m_name in FALLBACK_MODELS:
        if m_name not in model_list:
            model_list.append(m_name)

    training_label = st.radio(
        "🎓 วิธีนับชั่วโมงใบประกาศ/เกียรติบัตรอบรม:",
        ["ตามชั่วโมงที่ระบุในใบประกาศ (เช่น 3 ชม. = 3 ภาระงาน)",
         "ตามเกณฑ์งานพัฒนาตนเอง (1 ชม./วันทำการ)"],
        help="เกณฑ์ข้อ 12 ระบุ 1 ชม. ต่อ 1 วันทำการที่เข้าร่วมตามแผนพัฒนารายบุคคล โปรดตรวจสอบแนวปฏิบัติของคณะ/มหาวิทยาลัยว่าใช้แบบใด",
    )
    training_mode = "cert" if training_label.startswith("ตามชั่วโมง") else "rubric"

    if active_api_key:
        st.success("🟢 พบ Gemini API Key", icon="🔑")
    else:
        st.warning("🟡 ยังไม่ได้ใส่ Gemini API Key — ระบบจะอ่านเอกสารสแกนอัตโนมัติไม่ได้ แต่ยังกรอกข้อมูลเองได้ตามปกติ", icon="⚠️")

    if fitz is None:
        st.caption("ℹ️ ไม่พบ PyMuPDF: ระบบยังส่ง PDF เข้า Gemini ได้โดยตรง แต่จะไม่มีโหมดแปลงหน้าเป็นภาพสำรอง")

    if st.button("🧪 ทดสอบ Gemini API", use_container_width=True):
        if not active_api_key:
            st.warning("กรุณาใส่ Gemini API Key ก่อนทดสอบ")
        else:
            try:
                _, used_model = call_gemini(
                    active_api_key, model_list,
                    [{"text": 'ตอบกลับเป็น JSON เท่านั้น: {"status": "ok"}'}],
                )
                st.success("✅ เชื่อมต่อ Gemini สำเร็จ (รุ่น: " + used_model + ")")
            except GeminiError as ex:
                st.warning("⚠️ ทดสอบไม่สำเร็จ: " + str(ex))
            except Exception as ex:
                st.warning("⚠️ ทดสอบไม่สำเร็จ: " + str(ex))

    st.markdown("### ☁️ สถานะการเชื่อมต่อ")
    if use_gsheets:
        st.success("🟢 เชื่อมต่อ Google Sheets สำเร็จ", icon="☁️")
    else:
        st.info("🟡 โหมดบันทึกในระบบส่วนบุคคล (Local Session)", icon="💾")

    with st.expander("🧪 ทดสอบระบบ Google Sheets"):
        if st.button("กดทดสอบอ่าน-เขียน Google Sheets"):
            if use_gsheets and conn:
                try:
                    test_df = conn.read(ttl="0")
                    st.success("✅ อ่าน Google Sheets สำเร็จ! พบข้อมูลทั้งหมด " + str(len(test_df)) + " แถว")
                except Exception as ex:
                    st.error(f"❌ ไม่สามารถเชื่อมต่อได้: {ex}")
            else:
                st.warning("ยังไม่ได้เปิดใช้ Google Sheets ใน Secrets")

# ---------------------------------------------------------
# 8. Main Interface & Tabs
# ---------------------------------------------------------
st.markdown(f'<div class="main-header">🏛️ SKRU Academic Workload AI Assistant ({APP_VERSION})</div>', unsafe_allow_html=True)
effective_email = user_email.strip() if user_email and user_email.strip() != "" else "guest@skru.ac.th"
st.markdown(
    f'<div class="sub-header">ระบบช่วยสกัด เรียบเรียงภาษาทางการ และประเมินภาระงานตามเกณฑ์ มรภ.สงขลา (มติกช.) | ผู้ใช้: <b>{effective_email}</b></div>',
    unsafe_allow_html=True,
)

if st.session_state.get("show_saved_success"):
    st.markdown("""
    <div class="success-alert">
        🎉 บันทึกข้อมูลเข้าคลังภาระงานสะสมเรียบร้อยแล้ว! สามารถสลับไปดูตารางใน Tab 2 หรือดาวน์โหลดไฟล์ Word ใน Tab 3 ได้ทันที
    </div>
    """, unsafe_allow_html=True)
    st.session_state["show_saved_success"] = False

tab1, tab2, tab3 = st.tabs(["📥 1. สกัดและเรียบเรียงภาระงาน (AI)", "📊 2. คลังภาระงานสะสม", "📄 3. สรุปแบบ ป-มร.สข. 01 (Word)"])

all_data_df = load_all_data()

# ---------------------------------------------------------
# TAB 1
# ---------------------------------------------------------
with tab1:
    if not active_api_key:
        st.warning(
            "⚠️ ยังไม่ได้ระบุ Gemini API Key — ระบบจะอ่านไฟล์สแกน/รูปภาพด้วย AI ไม่ได้ "
            "กรุณากรอก API Key ในแถบด้านซ้าย (หรือตั้งค่า GEMINI_API_KEY ใน Secrets) "
            "หรือจะพิมพ์/กรอกข้อมูลในช่องด้านขวาด้วยตัวเองแล้วบันทึกก็ได้",
            icon="⚠️",
        )

    col_a, col_b = st.columns([1, 1], gap="large")

    with col_a:
        st.subheader("📝 ขั้นตอนที่ 1: ส่งไฟล์ / ภาพคำสั่ง / พิมพ์ข้อความ")

        category_input = st.selectbox("📂 เลือกหมวดภาระงานเบื้องต้น (หรือให้ AI ประเมิน):", [AUTO_CATEGORY] + CATEGORIES)

        input_method = st.radio("เลือกวิธีป้อนข้อมูลให้ AI:", [
            "📤 อัปโหลดไฟล์เอกสาร/คำสั่ง (PDF, PNG, JPG)",
            "📷 ถ่ายภาพคำสั่งจากกล้องมือถือ",
            "🔗 วางลิงก์ Google Drive (สำหรับมือถือที่เลือกไฟล์ไม่ผ่าน)",
            "✍️ พิมพ์รายละเอียดภาระงานเอง",
        ])

        uploaded_file = None
        raw_text_input = ""
        drive_url = ""

        if "อัปโหลด" in input_method:
            uploaded_file = st.file_uploader("แนบคำสั่ง/ประกาศ/วุฒิบัตร (PDF, JPG, PNG):", type=["pdf", "png", "jpg", "jpeg", "webp"])
        elif "ถ่ายภาพ" in input_method:
            uploaded_file = st.camera_input("ถ่ายภาพคำสั่งจากกล้องมือถือ")
        elif "Google Drive" in input_method:
            drive_url = st.text_input(
                "ลิงก์ไฟล์ Google Drive:",
                placeholder="https://drive.google.com/file/d/xxxx/view",
                help="ใน Drive กด ⋮ > แชร์ > ตั้ง 'ทุกคนที่มีลิงก์' แล้วคัดลอกลิงก์มาวาง (ควรปิดแชร์หลังใช้เสร็จ)",
            )
        else:
            raw_text_input = st.text_area(
                "พิมพ์รายละเอียดภาระงานหรือข้อความในคำสั่ง:",
                placeholder="เช่น ผ่านการอบรมหลักสูตรคณาจารย์นิเทศ CWIE โครงการ TCU กระทรวง อว. วันที่ 4 พ.ค. 2569...",
            )

        btn_ai_process = st.button("🤖 ให้ AI อ่านเอกสารและวิเคราะห์ตามเกณฑ์ มรภ.สงขลา", type="primary", use_container_width=True)

    if btn_ai_process:
        if uploaded_file is None and not raw_text_input.strip() and not drive_url.strip():
            st.session_state["ai_notice"] = ("warning", "⚠️ กรุณาอัปโหลดไฟล์ ถ่ายภาพ หรือพิมพ์ข้อความรายละเอียดก่อนกดปุ่มครับ")
        else:
            with st.spinner("กำลังอ่านเอกสาร (OCR/Vision) และประเมินภาระงานตามเกณฑ์ มรภ.สงขลา..."):
                result = None
                notice = None
                filename = ""
                local_text = ""
                file_bytes = b""
                mime = ""
                typed_text = ""

                garbled = False
                try:
                    if uploaded_file is not None:
                        file_bytes = uploaded_file.getvalue()
                        filename = getattr(uploaded_file, "name", "") or ""
                        mime = guess_mime(uploaded_file)
                        if filename.lower().endswith(".pdf"):
                            mime = "application/pdf"
                        if mime == "application/pdf":
                            local_text = extract_text_from_pdf_bytes(file_bytes)
                    elif drive_url.strip():
                        file_bytes, mime = fetch_drive_file(drive_url.strip())
                        if mime == "application/pdf":
                            local_text = extract_text_from_pdf_bytes(file_bytes)
                    else:
                        typed_text = raw_text_input.strip()
                        local_text = typed_text
                    if local_text and not typed_text and thai_text_looks_garbled(local_text):
                        garbled = True
                        local_text = ""
                except Exception as e:
                    notice = ("warning", "⚠️ อ่านไฟล์ไม่สำเร็จ: " + str(e) + " — กรุณากรอกข้อมูลด้วยตนเอง")

                is_scan_pdf = (mime == "application/pdf" and not local_text and not garbled)

                if notice is None:
                    if active_api_key:
                        try:
                            raw_json, used_model = cached_analyze(
                                file_bytes, mime, category_input, typed_text, tuple(model_list), active_api_key
                            )
                            result = finalize_result(raw_json, local_text, filename, category_input, training_mode)
                            msg = "✅ Gemini (" + used_model + ") อ่านเอกสารและร่างข้อความทางการเรียบร้อยแล้ว"
                            if is_scan_pdf:
                                msg += " (ตรวจพบว่าเป็น PDF สแกน ไม่มี text layer จึงอ่านจากภาพโดยตรง)"
                            if garbled:
                                msg += " (text layer ของ PDF สระ/วรรณยุกต์หาย จึงอ่านจากภาพแทน)"
                            if result["is_training"]:
                                msg += " | ตรวจพบเป็นใบประกาศ/เกียรติบัตรอบรม → จัดหมวด 5 และนับชั่วโมงตามโหมดที่เลือกในแถบด้านซ้าย"
                            if result["rule_applied"]:
                                msg += " | ใช้กฎงานสร้างสรรค์นานาชาติที่ได้รับรางวัล/ดีเยี่ยม = 24.0 ภาระงาน"
                            if result["notes"]:
                                msg += " | ข้อสังเกต: " + result["notes"]
                            notice = ("success", msg)
                        except GeminiAuthError as e:
                            notice = ("warning", "⚠️ " + str(e) + " — กรุณาตรวจสอบ API Key หรือกรอกข้อมูลด้วยตนเอง")
                        except GeminiError as e:
                            notice = ("warning", "⚠️ Gemini ขัดข้องชั่วคราว: " + str(e) + " — ใช้โหมดสำรองแล้ว คุณกรอก/แก้ข้อมูลเองได้ทุกช่อง")
                        except Exception as e:
                            notice = ("warning", "⚠️ เกิดข้อผิดพลาดที่ไม่คาดคิดขณะเรียก Gemini: " + str(e) + " — ใช้โหมดสำรองแทน")
                    else:
                        notice = ("warning", "⚠️ ยังไม่ได้ใส่ Gemini API Key จึงไม่สามารถให้ AI อ่านเอกสารได้ กรุณากรอก API Key ในแถบด้านซ้าย หรือกรอกข้อมูลในช่องด้านขวาด้วยตัวเอง")

                if result is None:
                    # โหมดสำรอง: ใช้ regex กับ text layer เท่าที่มี ไม่ใส่ข้อมูลสมมติ ไม่ใช้ชื่อไฟล์
                    try:
                        result = parse_text_local(local_text, filename, category_input, training_mode) if local_text else blank_result(category_input)
                    except Exception:
                        result = blank_result(category_input)
                    if is_scan_pdf or (mime.startswith("image/") and mime != ""):
                        notice = (notice[0], notice[1] + " | เอกสารนี้เป็นภาพ/สแกน ต้องใช้ Gemini จึงจะอ่านได้ ช่องด้านขวาจึงว่างไว้ให้กรอกเอง")

                st.session_state["ai_notice"] = notice
                st.session_state["edit_cat"] = result["category"]
                st.session_state["edit_title"] = result["title"]
                st.session_state["edit_venue"] = result["venue"]
                st.session_state["edit_date"] = result["date"]
                st.session_state["edit_ref"] = result["ref"]
                st.session_state["edit_formula"] = result["formula"]
                st.session_state["edit_hours"] = float(result["hours"])
                st.session_state["auto_compose"] = True

    with col_b:
        st.subheader("🤖 ขั้นตอนที่ 2: ผลการประเมินจาก AI (ตรวจทาน/ปรับแก้ได้ทุกช่อง)")

        notice = st.session_state.get("ai_notice")
        if notice:
            if notice[0] == "success":
                st.success(notice[1])
            else:
                st.warning(notice[1])
        else:
            st.markdown(
                "<div class='info-alert'>💡 <b>อัปโหลดไฟล์แล้วกดปุ่ม '🤖 ให้ AI อ่านเอกสาร...'</b> "
                "ระบบจะสกัดเนื้อหาจริงและร่างข้อความทางการให้ คุณสามารถพิมพ์แก้ไขทุกช่องได้ทันที</div>",
                unsafe_allow_html=True,
            )

        defaults = {
            "edit_cat": CATEGORIES[1], "edit_title": "", "edit_venue": "", "edit_date": "",
            "edit_ref": "", "edit_formula": "", "edit_hours": 0.0, "auto_compose": True,
        }
        for k, v in defaults.items():
            if k not in st.session_state:
                st.session_state[k] = v
        if st.session_state["edit_cat"] not in CATEGORIES:
            st.session_state["edit_cat"] = CATEGORIES[1]

        final_cat = st.selectbox("📂 หมวดภาระงาน:", CATEGORIES, key="edit_cat")

        c_f1, c_f2 = st.columns(2)
        with c_f1:
            final_title = st.text_input("1. ชื่อบทบาท / โครงการ / ผลงาน:", key="edit_title")
            final_date = st.text_input("3. วันที่ปฏิบัติงาน:", key="edit_date")
            final_formula = st.text_input("5. สูตรคำนวณ (ในวงเล็บ):", key="edit_formula")
        with c_f2:
            final_venue = st.text_input("2. สถานที่จัด / หน่วยงาน:", key="edit_venue")
            final_ref = st.text_input("4. เลขที่คำสั่ง / หนังสืออ้างอิง:", key="edit_ref")
            final_hours = st.number_input("6. สรุปชั่วโมงภาระงานสุทธิ:", min_value=0.0, step=0.5, format="%.2f", key="edit_hours")

        auto_compose = st.checkbox("สร้างข้อความทางการอัตโนมัติจากช่องด้านบน (ปิดเพื่อพิมพ์แก้ข้อความเองทั้งหมด)", key="auto_compose")
        if auto_compose:
            st.session_state["edit_formal_text"] = compose_formal_text(
                final_title, final_venue, final_date, final_ref, final_formula, float(final_hours)
            )
        elif "edit_formal_text" not in st.session_state:
            st.session_state["edit_formal_text"] = ""

        final_formal_text = st.text_area(
            "📝 ข้อความภาษาทางการฉบับสมบูรณ์ (ที่จะบันทึกลงตารางและไฟล์ Word):",
            key="edit_formal_text",
            height=120,
        )

        st.markdown(f"""
        <div class="formatted-preview">
            <b>📂 หมวดงาน:</b> {final_cat}<br>
            <b>📄 อ้างอิง:</b> {final_ref}<br>
            <b>📊 ภาระงานสุทธิ:</b> <span style="color:green; font-size:1.2rem; font-weight:bold;">{float(final_hours):.1f} ภาระงาน</span>
        </div>
        """, unsafe_allow_html=True)

        btn_save_item = st.button("💾 บันทึกลงคลังภาระงานสะสม (ซิงค์ Cloud)", type="primary", use_container_width=True)

        if btn_save_item:
            if not final_formal_text.strip():
                st.warning("⚠️ ยังไม่มีข้อความภาระงานให้บันทึก กรุณากรอกข้อมูลก่อนครับ")
            else:
                new_entry = {
                    "email": effective_email,
                    "หมวดงาน": final_cat,
                    "รายการภาระงาน": final_formal_text,
                    "เลขคำสั่ง_อ้างอิง": final_ref,
                    "วันที่": final_date,
                    "ภาระงาน_ชม": float(final_hours),
                    "วันที่บันทึก": datetime.now().strftime("%Y-%m-%d %H:%M"),
                }
                updated_all_df = pd.concat([all_data_df, pd.DataFrame([new_entry])], ignore_index=True)
                success, msg = save_all_data(updated_all_df)

                if success:
                    st.balloons()
                    st.toast("🎉 บันทึกข้อมูลเข้าคลังภาระงานเรียบร้อยแล้ว!", icon="✅")
                    st.success(f"🎉 บันทึกรายการภาระงานสำหรับผู้ใช้ '{effective_email}' ลงคลังสะสมเรียบร้อยแล้ว!")
                    st.session_state["show_saved_success"] = True
                else:
                    st.error(f"เกิดข้อผิดพลาดในการบันทึกขึ้น Google Sheets (เก็บสำเนาในเครื่องไว้ให้แล้ว): {msg}")

# ---------------------------------------------------------
# TAB 2
# ---------------------------------------------------------
with tab2:
    st.subheader(f"📊 คลังภาระงานสะสมของผู้ใช้: {effective_email}")

    current_all_df = load_all_data()
    user_df = current_all_df[current_all_df["email"] == effective_email].copy()

    if user_df.empty:
        st.info(f"👋 ขณะนี้คลังภาระงานของผู้ใช้ '{effective_email}' ยังว่างเปล่า สามารถเริ่มบันทึกรายการแรกได้ใน Tab 1 ครับ")
    else:
        st.markdown("💡 **อาจารย์สามารถแก้ไขข้อความ หรือลบรายการในตารางได้โดยตรง แล้วกดปุ่มบันทึกด้านล่าง:**")

        display_cols = ["หมวดงาน", "รายการภาระงาน", "เลขคำสั่ง_อ้างอิง", "วันที่", "ภาระงาน_ชม"]

        edited_user_df = st.data_editor(
            user_df[display_cols].reset_index(drop=True),
            use_container_width=True,
            num_rows="dynamic",
            key="user_data_editor_v14",
        )

        col_m1, col_m2, col_m3 = st.columns([1.5, 1, 1])
        with col_m1:
            total_user_hours = pd.to_numeric(edited_user_df["ภาระงาน_ชม"], errors="coerce").fillna(0.0).sum()
            st.metric(label="📈 รวมภาระงานสะสมสุทธิ", value=f"{total_user_hours:.1f} ภาระงาน (ชั่วโมง)")

        with col_m2:
            if st.button("🔄 บันทึกการแก้ไขลงฐานข้อมูล", type="primary", use_container_width=True):
                other_users_df = current_all_df[current_all_df["email"] != effective_email]
                to_save = edited_user_df.copy()
                to_save["email"] = effective_email
                to_save["วันที่บันทึก"] = datetime.now().strftime("%Y-%m-%d %H:%M")
                to_save = ensure_schema(to_save)
                to_save = to_save[to_save["รายการภาระงาน"].str.strip() != ""]

                new_all_df = pd.concat([other_users_df, to_save], ignore_index=True)
                ok, err = save_all_data(new_all_df)
                if ok:
                    st.success("บันทึกการปรับเปลี่ยนลงฐานข้อมูลเรียบร้อยแล้ว!")
                    st.rerun()
                else:
                    st.error(f"ไม่สามารถบันทึกได้: {err}")

        with col_m3:
            confirm_delete = st.checkbox("ยืนยันการลบ", key="confirm_delete_all")
            if st.button("🗑️ ลบข้อมูลทั้งหมดของคุณ", type="secondary", use_container_width=True, disabled=not confirm_delete):
                other_users_df = current_all_df[current_all_df["email"] != effective_email]
                save_all_data(other_users_df)
                st.warning("เคลียร์คลังภาระงานของคุณเรียบร้อยแล้ว!")
                st.rerun()

# ---------------------------------------------------------
# TAB 3
# ---------------------------------------------------------
with tab3:
    st.subheader(f"📋 สรุปผลการประเมินภาระงานและสร้างไฟล์ Word (ผู้ใช้: {effective_email})")

    current_all_df = load_all_data()
    user_df = current_all_df[current_all_df["email"] == effective_email].copy()

    total_actual_hours = float(pd.to_numeric(user_df["ภาระงาน_ชม"], errors="coerce").fillna(0.0).sum()) if not user_df.empty else 0.0

    target_benchmark = 35.0
    achieved_ratio = min(total_actual_hours / target_benchmark, 1.0) if target_benchmark > 0 else 0.0
    score_component1 = achieved_ratio * 70.0

    st.markdown("### 📊 คะแนนประเมินรวมรายองค์ประกอบ")

    c1, c2, c3 = st.columns(3)
    with c1:
        st.metric("องค์ประกอบที่ 1 (ผลสัมฤทธิ์ภาระงาน)", f"{score_component1:.2f} / 70 คะแนน", delta=f"ภาระงานสะสม {total_actual_hours:.1f} ชม.")
    with c2:
        competency_score = st.number_input("องค์ประกอบที่ 2 (สมรรถนะ/พฤติกรรม):", min_value=0.0, max_value=30.0, value=30.0, step=0.5)
    with c3:
        total_score = score_component1 + competency_score
        grade_text = "ต้องปรับปรุง (< 60)"
        if total_score >= 90:
            grade_text = "ดีเด่น (90.00 - 100)"
        elif total_score >= 80:
            grade_text = "ดีมาก (80.00 - 89.99)"
        elif total_score >= 70:
            grade_text = "ดี (70.00 - 79.99)"
        elif total_score >= 60:
            grade_text = "พอใช้ (60.00 - 69.99)"

        st.metric("คะแนนประเมินรวมสุทธิ", f"{total_score:.2f} / 100 คะแนน", delta=grade_text)

    st.divider()

    def generate_docx():
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        buffer = io.BytesIO()
        doc = Document()

        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = title_p.add_run("แบบสรุปการประเมินผลการปฏิบัติราชการของบุคลากรสายวิชาการ (ป-มร.สข. 01)")
        run_title.bold = True
        run_title.font.size = Pt(16)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.add_run(f'มหาวิทยาลัยราชภัฏสงขลา | รอบการประเมิน {datetime.now().strftime("%Y")}')

        doc.add_paragraph(f"ผู้รับการประเมิน: {effective_email}")
        doc.add_paragraph(f"ตำแหน่ง/สังกัด: {role}")
        doc.add_paragraph(f"กลุ่มการประเมิน: {track_type}")
        doc.add_paragraph(f'วันที่ออกรายงาน: {datetime.now().strftime("%d/%m/%Y")}')

        doc.add_heading("1. สรุปคะแนนการประเมินผลการปฏิบัติงาน", level=2)
        doc.add_paragraph(f"• องค์ประกอบที่ 1 (ผลสัมฤทธิ์ของงาน): {score_component1:.2f} / 70 คะแนน")
        doc.add_paragraph(f"• องค์ประกอบที่ 2 (พฤติกรรม/สมรรถนะ): {competency_score:.2f} / 30 คะแนน")
        doc.add_paragraph(f"• คะแนนรวมสุทธิ: {total_score:.2f} / 100 คะแนน (ระดับผลการประเมิน: {grade_text})")

        doc.add_heading("2. รายละเอียดภาระงานสะสม (เรียบเรียงภาษาทางการ)", level=2)

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "หมวดงาน"
        hdr_cells[1].text = "รายการภาระงานและสูตรการคำนวณ (ภาษาทางการ)"
        hdr_cells[2].text = "เลขที่คำสั่ง/อ้างอิง"
        hdr_cells[3].text = "ภาระงาน (ชม.)"

        if not user_df.empty:
            for _, row in user_df.iterrows():
                row_cells = table.add_row().cells
                row_cells[0].text = str(row.get("หมวดงาน", ""))
                row_cells[1].text = str(row.get("รายการภาระงาน", ""))
                row_cells[2].text = str(row.get("เลขคำสั่ง_อ้างอิง", ""))
                row_cells[3].text = f"{to_float(row.get('ภาระงาน_ชม', 0)):.1f}"

        doc.add_paragraph(f"\nรวมภาระงานสะสมสุทธิทั้งสิ้น: {total_actual_hours:.1f} ภาระงาน (ชั่วโมง)")
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    try:
        docx_file = generate_docx()
        file_email_slug = effective_email.split("@")[0]
        st.download_button(
            label="📥 ดาวน์โหลดแบบสรุป ป-มร.สข. 01 (.docx)",
            data=docx_file,
            file_name=f"ป-มร.สข.01_{file_email_slug}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
        )
    except ImportError:
        st.warning("⚠️ ยังไม่ได้ติดตั้ง python-docx จึงสร้างไฟล์ Word ไม่ได้ (เพิ่ม python-docx ใน requirements.txt)")
    except Exception as e:
        st.warning(f"⚠️ สร้างไฟล์ Word ไม่สำเร็จ: {e}")
